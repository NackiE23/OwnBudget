"""
Імпорт історії з Excel-таблиці бюджету (етап 3 ТЗ).

Таблиця: один рядок на місяць. Лівий блок (Month … Crypto) знаходиться за
назвами в рядку заголовків, правий блок категорій за назвами над ним: кожна
категорія займає два стовпці, «Як Є» і «Ціль/Межа». Значення беруться з кешу
формул, а самі формули дають відсоток ліміту й вираз Salary.
"""

import datetime
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

import openpyxl
from django.db import transaction

from budget import calculations, expressions, services
from budget.models import (
    AllocationCategory,
    Currency,
    LineType,
    Month,
    MonthAllocation,
    MonthLine,
    SavingsDestination,
    SavingsEntry,
)
from budget.parsing import ParsedEntry, parse_savings_line, parse_terms

PLACEHOLDERS = {"", "-", "—", "–"}
PERIOD_RE = re.compile(r"^(\d{2})\.(\d{2})$")
# Ліміт як частка іншої комірки: =G13*30% або =N14*20%.
SHARE_FORMULA_RE = re.compile(r"^=\s*\$?[A-Z]+\$?\d+\s*\*\s*(\d+(?:[.,]\d+)?)\s*%$")
COLUMN_MONTH = "Month"
COLUMN_TOTAL = "Total expenses"
COLUMN_SALARY = "Salary"
COLUMN_FOOD = "Food"
COLUMN_SAVINGS = "Savings"
COLUMN_INCOMES = "Additional incomes"
COLUMN_NOTES = "Notes"
COLUMN_CRYPTO = "Crypto"
EXPENSE_COLUMNS = ("Rent", "Utility services", "Taxes")
REQUIRED_COLUMNS = (COLUMN_MONTH, COLUMN_SALARY)
FOOD_CATEGORY = "Їжа"
CRYPTO_DESTINATION = "Crypto"


class WorkbookError(Exception):
    """Файл не схожий на таблицю бюджету."""


class _DryRun(Exception):
    pass


@dataclass
class Report:
    dry_run: bool = False
    created: list[str] = field(default_factory=list)
    updated: list[str] = field(default_factory=list)
    skipped_closed: list[str] = field(default_factory=list)
    empty_rows: list[str] = field(default_factory=list)
    closed: list[str] = field(default_factory=list)
    new_destinations: list[str] = field(default_factory=list)
    warnings: list[tuple[str, str]] = field(default_factory=list)

    def warn(self, label: str, message: str) -> None:
        self.warnings.append((label, message))

    def lines(self) -> list[str]:
        out = []
        if self.dry_run:
            out.append("ПРОБНИЙ ЗАПУСК: зміни не збережено.")
        out.append(f"Створено місяців: {len(self.created)}. Оновлено: {len(self.updated)}.")
        if self.updated:
            out.append(f"  Оновлено: {', '.join(self.updated)}")
        if self.closed:
            out.append(f"Закрито: {len(self.closed)} (усі, крім останнього місяця файлу).")
        if self.skipped_closed:
            out.append(
                f"Пропущено закриті місяці ({len(self.skipped_closed)}), щоб перезаписати, додай --overwrite-closed: "
                + ", ".join(self.skipped_closed)
            )
        if self.empty_rows:
            out.append(f"Порожні рядки без даних пропущено: {', '.join(self.empty_rows)}")
        if self.new_destinations:
            out.append(f"Нові призначення накопичень: {', '.join(self.new_destinations)}")
        if self.warnings:
            out.append(f"Перевір вручну ({len(self.warnings)}):")
            out.extend(f"  {label}: {message}" for label, message in self.warnings)
        return out


# --- розбір значень -------------------------------------------------------


def is_placeholder(value) -> bool:
    return value is None or (isinstance(value, str) and value.strip() in PLACEHOLDERS)


def to_decimal(value) -> Decimal | None:
    if is_placeholder(value):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    try:
        return Decimal(str(value).strip().replace(" ", "").replace(" ", "").replace(",", "."))
    except InvalidOperation:
        return None


def parse_period(value) -> datetime.date | None:
    match = PERIOD_RE.match(str(value or "").strip())
    if not match:
        return None
    month, year = int(match[1]), int(match[2])
    if not 1 <= month <= 12:
        return None
    return datetime.date(2000 + year, month, 1)


def text_lines(value) -> list[str]:
    if is_placeholder(value):
        return []
    if isinstance(value, (int, float)):
        return [format(Decimal(str(value)).normalize(), "f")]
    return [line.strip() for line in str(value).splitlines() if line.strip() not in PLACEHOLDERS]


def expression(formula) -> str:
    """
    Формула з обчисленням («=1259+200») у вигляді, як її зберігає застосунок.
    Посилання на комірки («=D17+E17») чи одне число («=1233») не зберігаються.
    """
    if not isinstance(formula, str) or not formula.startswith("=") or not expressions.is_expression(formula):
        return ""
    return expressions.parse_amount(formula)[1]


def money_text(value: Decimal) -> str:
    return f"{calculations.money(value):,.2f}".replace(",", " ")


# --- читання аркуша -------------------------------------------------------


@dataclass
class Sheet:
    formulas: object  # аркуш із формулами та заливкою
    values: object  # той самий аркуш із кешованими значеннями
    header_row: int
    columns: dict[str, int]
    categories: list[tuple[str, int]]  # (назва, стовпець «Як Є»); «Ціль/Межа» праворуч

    def value(self, row: int, column: int | None):
        return self.values.cell(row, column).value if column else None

    def formula(self, row: int, column: int | None):
        return self.formulas.cell(row, column).value if column else None


def open_sheet(path, sheet_name: str | None = None) -> Sheet:
    formulas_book = openpyxl.load_workbook(path)
    values_book = openpyxl.load_workbook(path, data_only=True)
    formulas = formulas_book[sheet_name] if sheet_name else formulas_book.active
    return build_sheet(formulas, values_book[formulas.title])


def build_sheet(formulas, values) -> Sheet:
    """Знаходить заголовки. formulas і values це той самий аркуш, прочитаний двічі."""
    header_row = None
    for row in formulas.iter_rows(max_row=min(formulas.max_row, 50)):
        if any(isinstance(cell.value, str) and cell.value.strip() == COLUMN_MONTH for cell in row):
            header_row = row[0].row
            break
    if header_row is None:
        raise WorkbookError(f"Не знайдено рядок заголовків зі стовпцем «{COLUMN_MONTH}».")

    columns = {}
    for cell in formulas[header_row]:
        if isinstance(cell.value, str) and cell.value.strip():
            columns.setdefault(cell.value.strip(), cell.column)
    missing = [name for name in REQUIRED_COLUMNS if name not in columns]
    if missing:
        raise WorkbookError(f"Немає стовпців: {', '.join(missing)}.")

    left_edge = max(columns[name] for name in columns if name in {COLUMN_CRYPTO, COLUMN_NOTES, COLUMN_MONTH})
    categories = []
    if header_row > 1:
        for cell in formulas[header_row - 1]:
            if cell.column > left_edge and isinstance(cell.value, str) and cell.value.strip():
                categories.append((cell.value.strip(), cell.column))
    return Sheet(formulas, values, header_row, columns, categories)


# --- імпорт ---------------------------------------------------------------


class Importer:
    def __init__(self, sheet: Sheet, *, overwrite_closed: bool = False):
        self.sheet = sheet
        self.overwrite_closed = overwrite_closed
        self.report = Report()
        self.categories = {c.name: c for c in AllocationCategory.objects.all()}
        self.line_types = {t.name: t for t in LineType.objects.all()}
        self.destinations = {d.name.lower(): d for d in SavingsDestination.objects.all()}

        unknown = [name for name, _ in sheet.categories if name not in self.categories]
        if unknown:
            raise WorkbookError(f"Невідомі категорії в таблиці: {', '.join(unknown)}. Додай їх у довідник.")
        absent = [name for name in (*EXPENSE_COLUMNS, COLUMN_INCOMES) if name in sheet.columns and name not in self.line_types]
        if absent:
            raise WorkbookError(f"Немає типів рядків: {', '.join(absent)}. Додай їх у довідник.")

    def run(self, *, dry_run: bool = False) -> Report:
        self.report.dry_run = dry_run
        try:
            with transaction.atomic():
                self._import_rows()
                if dry_run:
                    raise _DryRun
        except _DryRun:
            pass
        return self.report

    def _rows(self) -> list[tuple[datetime.date, int]]:
        rows = []
        month_column = self.sheet.columns[COLUMN_MONTH]
        for row in range(self.sheet.header_row + 1, self.sheet.formulas.max_row + 1):
            period = parse_period(self.sheet.value(row, month_column))
            if period is None:
                continue
            if self._row_is_empty(row, month_column):
                self.report.empty_rows.append(period.strftime("%m.%y"))
                continue
            rows.append((period, row))
        return sorted(rows)

    def _row_is_empty(self, row: int, month_column: int) -> bool:
        last_column = max([*self.sheet.columns.values(), *(column + 1 for _, column in self.sheet.categories)])
        return all(
            is_placeholder(self.sheet.value(row, column))
            for column in range(1, last_column + 1)
            if column != month_column
        )

    def _import_rows(self) -> None:
        rows = self._rows()
        imported = []
        for period, row in rows:
            month = self._import_month(period, row)
            if month is not None:
                imported.append(month)
        if not rows:
            return
        latest_period = rows[-1][0]
        for month in imported:
            if month.period != latest_period:
                services.close_month(month)
                self.report.closed.append(str(month))

    def _import_month(self, period: datetime.date, row: int) -> Month | None:
        label = period.strftime("%m.%y")
        month = Month.objects.filter(period=period).first()
        if month and month.is_closed:
            if not self.overwrite_closed:
                self.report.skipped_closed.append(label)
                return None
            services.reopen_month(month)

        if month:
            # Масове видалення: рядки місяця замінюються вмістом файлу.
            month.allocations.all().delete()
            month.lines.all().delete()
            month.savings_entries.all().delete()
            self.report.updated.append(label)
        else:
            month = Month(period=period)
            self.report.created.append(label)

        self._fill_month(month, row, label)
        month.save()
        excel_limits = self._import_allocations(month, row, label)
        self._import_lines(month, row, label)
        self._import_savings(month, row, label)

        month.recalculate()
        self._compare_limits(month, excel_limits, label)
        return month

    def _fill_month(self, month: Month, row: int, label: str) -> None:
        columns = self.sheet.columns
        salary_column = columns[COLUMN_SALARY]
        salary = to_decimal(self.sheet.value(row, salary_column))
        if salary is None and not is_placeholder(self.sheet.value(row, salary_column)):
            self.report.warn(label, f"Salary «{self.sheet.value(row, salary_column)}» не число, записано 0.")
        month.salary = calculations.money(salary or Decimal("0"))
        month.salary_expression = expression(self.sheet.formula(row, salary_column))

        total = self.sheet.value(row, columns.get(COLUMN_TOTAL))
        month.total_expenses = calculations.money(to_decimal(total)) if to_decimal(total) is not None else None
        notes = self.sheet.value(row, columns.get(COLUMN_NOTES))
        month.notes = "\n".join(line.rstrip() for line in str(notes).strip().splitlines()) if not is_placeholder(notes) else ""

    def _import_allocations(self, month: Month, row: int, label: str) -> dict[str, Decimal]:
        """Створює розподіл за категоріями; повертає ліміти з таблиці для порівняння."""
        food = to_decimal(self.sheet.value(row, self.sheet.columns.get(COLUMN_FOOD)))
        excel_limits = {}
        food_used = False
        food_expression = expression(self.sheet.formula(row, self.sheet.columns.get(COLUMN_FOOD)))
        for name, column in self.sheet.categories:
            category = self.categories[name]
            actual = to_decimal(self.sheet.value(row, column))
            actual_expression = expression(self.sheet.formula(row, column))
            percent, limit_amount = self._limit(row, column + 1)

            if name == FOOD_CATEGORY and food is not None:
                food_used = True
                if actual is None:
                    actual, actual_expression = food, food_expression
                elif calculations.money(actual) != calculations.money(food):
                    self.report.warn(
                        label,
                        f"{name}: «Як Є» {money_text(actual)} і Food {money_text(food)} різні, взято «Як Є».",
                    )

            if actual is None and percent is None and limit_amount is None:
                continue
            if limit_amount is not None:
                excel_limits[name] = limit_amount
            MonthAllocation.objects.create(
                month=month,
                category=category,
                percent=percent,
                limit_amount=calculations.money(limit_amount) if limit_amount is not None else None,
                actual=calculations.money(actual) if actual is not None else None,
                actual_expression=actual_expression,
            )

        if food is not None and not food_used:
            # Старі місяці без блоку категорій: Food стає фактичною сумою «Їжа».
            MonthAllocation.objects.create(
                month=month,
                category=self.categories[FOOD_CATEGORY],
                actual=calculations.money(food),
                actual_expression=food_expression,
            )
        return excel_limits

    def _limit(self, row: int, column: int) -> tuple[Decimal | None, Decimal | None]:
        """(відсоток, ліміт) з комірки «Ціль/Межа»."""
        formula = self.sheet.formula(row, column)
        value = to_decimal(self.sheet.value(row, column))
        if isinstance(formula, str):
            match = SHARE_FORMULA_RE.match(formula.replace(" ", ""))
            if match:
                return Decimal(match[1].replace(",", ".")), value
        if value is not None and 0 < value <= 1:
            # У частині місяців замість ліміту записано частку: 0.3 означає 30%.
            return calculations.money(value * calculations.HUNDRED), None
        return None, value

    def _compare_limits(self, month: Month, excel_limits: dict[str, Decimal], label: str) -> None:
        for allocation in month.allocations.select_related("category"):
            expected = excel_limits.get(allocation.category.name)
            if allocation.percent is None or expected is None or allocation.limit_amount is None:
                continue
            if abs(allocation.limit_amount - calculations.money(expected)) > Decimal("0.01"):
                self.report.warn(
                    label,
                    f"{allocation.category.name}: у таблиці ліміт {money_text(expected)}, "
                    f"за правилами {money_text(allocation.limit_amount)}.",
                )

    def _import_lines(self, month: Month, row: int, label: str) -> None:
        columns = self.sheet.columns
        for name in EXPENSE_COLUMNS:
            column = columns.get(name)
            raw = self.sheet.value(row, column)
            amount = to_decimal(raw)
            if amount is None:
                if not is_placeholder(raw):
                    self.report.warn(label, f"{name}: «{raw}» не число, рядок пропущено.")
                continue
            MonthLine.objects.create(
                month=month,
                line_type=self.line_types[name],
                amount=calculations.money(amount),
                amount_expression=expression(self.sheet.formula(row, column))[:255],
            )

        raw = self.sheet.value(row, columns.get(COLUMN_INCOMES))
        lines = text_lines(raw)
        if not lines:
            return
        total = Decimal("0")
        for line in lines:
            terms = parse_terms(line)
            if terms is None:
                self.report.warn(label, f"Additional incomes: рядок «{line}» не розібрано, у суму не входить.")
                continue
            for term in terms:
                if term.currency == Currency.UAH:
                    total += term.amount
                else:
                    self.report.warn(
                        label,
                        f"Additional incomes: {term.amount} {term.currency} з «{line}» не входить у суму (лише гривні).",
                    )
        note = "; ".join(lines)
        if len(note) > 255:
            self.report.warn(label, "Additional incomes: текст довший за 255 символів, обрізано в примітці.")
        MonthLine.objects.create(
            month=month, line_type=self.line_types[COLUMN_INCOMES], amount=calculations.money(total), note=note[:255]
        )

    def _import_savings(self, month: Month, row: int, label: str) -> None:
        columns = self.sheet.columns
        for line in text_lines(self.sheet.value(row, columns.get(COLUMN_SAVINGS))):
            for entry in parse_savings_line(line):
                self._create_entry(month, entry, label, column=COLUMN_SAVINGS, destination=entry.destination)

        for line in text_lines(self.sheet.value(row, columns.get(COLUMN_CRYPTO))):
            # «2000 на USDT»: гривні, переведені в крипту. Текст лишається приміткою.
            for entry in parse_savings_line(line):
                self._create_entry(
                    month, entry, label, column=COLUMN_CRYPTO, destination=CRYPTO_DESTINATION, note=entry.destination
                )

    def _create_entry(
        self, month: Month, entry: ParsedEntry, label: str, *, column: str, destination: str, note: str = ""
    ) -> None:
        if entry.amount is None:
            self.report.warn(label, f"{column}: «{entry.raw_text}» складний, записано без суми.")
        SavingsEntry.objects.create(
            month=month,
            amount=calculations.money(entry.amount) if entry.amount is not None else None,
            currency=entry.currency,
            destination=self._destination(destination) if destination else None,
            raw_text=entry.raw_text[:500],
            note=note[:255],
        )

    def _destination(self, name: str) -> SavingsDestination:
        key = name.lower()
        if key not in self.destinations:
            self.destinations[key] = SavingsDestination.objects.create(name=name[:100])
            self.report.new_destinations.append(name)
        return self.destinations[key]
