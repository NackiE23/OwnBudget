import datetime
import tempfile
from decimal import Decimal
from io import StringIO
from pathlib import Path

import openpyxl
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import SimpleTestCase, TestCase

from budget.importer import Importer, build_sheet
from budget.models import Month, SavingsEntry
from budget.parsing import parse_savings_line, parse_terms

D = Decimal
AUG = datetime.date(2025, 8, 1)
SEP = datetime.date(2025, 9, 1)
OCT = datetime.date(2025, 10, 1)

HEADERS = {
    "B": "Month",
    "C": "Total expenses",
    "D": "Rent",
    "E": "Utility services",
    "F": "Taxes",
    "G": "Salary",
    "H": "Food",
    "I": "Savings",
    "J": "Additional incomes",
    "K": "Notes",
    "L": "Crypto",
    "N": "Як Є",
    "O": "Ціль",
    "P": "Як Є",
    "Q": "Межа",
    "R": "Як Є",
    "S": "Межа",
}
CATEGORIES = {"N": "Загальні Накопичення", "P": "Їжа", "R": "Знято з накопичень"}


class FakeWorkbook:
    """
    Синтетична таблиця в розкладці справжньої. openpyxl не обчислює формули,
    тож аркуш із формулами й аркуш зі значеннями заповнюються окремо.
    """

    def __init__(self):
        self.formulas = openpyxl.Workbook().active
        self.values = openpyxl.Workbook().active
        self.set("B1", "Бюджет")
        for column, name in CATEGORIES.items():
            self.set(f"{column}3", name)
        for column, name in HEADERS.items():
            self.set(f"{column}4", name)

    def set(self, coordinate: str, value, formula: str | None = None) -> None:
        self.formulas[coordinate] = formula if formula is not None else value
        self.values[coordinate] = value

    def sheet(self):
        return build_sheet(self.formulas, self.values)

    def save_values(self, path: Path) -> None:
        self.values.parent.save(path)


def sample_workbook() -> FakeWorkbook:
    book = FakeWorkbook()
    # Майбутній місяць без даних.
    book.set("B5", "11.25")

    # Поточний місяць: формули лімітів, Savings із різними форматами.
    book.set("B6", "10.25")
    book.set("C6", 72000)
    book.set("D6", 15000)
    book.set("E6", 1459, "=1259+200")
    book.set("F6", "-")
    book.set("G6", 50000, "=1000+49000")
    book.set("I6", "$100 моно\n2000+4000 облігації\n8000 + $200\n4000 (-1500)")
    book.set("J6", "3500 + 100€ подарунок\n1100 Петро + Іра")
    book.set("K6", "-2000 лампа  \n-500 книги")
    book.set("L6", "2000 на USDT")
    book.set("N6", 12000)
    book.set("O6", 15000, "=G6*30%")
    book.set("P6", "-")
    book.set("Q6", 10000, "=G6*20%")
    book.set("S6", 2000, "=N7*20%")

    # Перший місяць нового блоку: частки замість лімітів, Food і «Як Є» різні.
    book.set("B7", "09.25")
    book.set("G7", 40000)
    book.set("H7", 9000)
    book.set("N7", 10000)
    book.set("O7", 0.3)
    book.set("P7", 8500)
    book.set("Q7", 0.2)

    # Старий місяць: лише лівий блок.
    book.set("B8", "08.25")
    book.set("D8", "—")
    book.set("G8", 30000)
    book.set("H8", 7000)
    book.set("I8", 2000)

    book.set("B9", "MAX:")
    book.set("C9", 72000, "=MAX(C5:C8)")
    return book


def run(book: FakeWorkbook, **options):
    dry_run = options.pop("dry_run", False)
    return Importer(book.sheet(), **options).run(dry_run=dry_run)


class ParseTermsTests(SimpleTestCase):
    def test_currencies(self):
        self.assertEqual(
            [(t.amount, t.currency, t.text) for t in parse_terms("$100 моно + 100€ подарунок + 200zl")],
            [(D("100"), "USD", "моно"), (D("100"), "EUR", "подарунок"), (D("200"), "PLN", "")],
        )

    def test_plus_inside_text_is_one_term(self):
        terms = parse_terms("1100 Петро + Іра")
        self.assertEqual([(t.amount, t.text) for t in terms], [(D("1100"), "Петро + Іра")])

    def test_complex_lines(self):
        for line in ["4000 (-1500)", "(-100$ | -4000)", "-1500 переїзд", "*було 2000 раніше"]:
            self.assertIsNone(parse_terms(line), line)

    def test_savings_sum_with_shared_destination(self):
        [entry] = parse_savings_line("2000+4000 облігації")
        self.assertEqual((entry.amount, entry.currency, entry.destination), (D("6000"), "UAH", "облігації"))

    def test_savings_mixed_currencies_split(self):
        entries = parse_savings_line("8000 + $200")
        self.assertEqual([(e.amount, e.currency) for e in entries], [(D("8000"), "UAH"), (D("200"), "USD")])
        self.assertEqual({e.raw_text for e in entries}, {"8000 + $200"})


class ImportTests(TestCase):
    def test_imports_months_from_sheet(self):
        report = run(sample_workbook())

        self.assertEqual(list(Month.objects.order_by("period").values_list("period", flat=True)), [AUG, SEP, OCT])
        self.assertEqual(report.empty_rows, ["11.25"])
        self.assertEqual(report.closed, ["08.25", "09.25"])
        self.assertEqual(Month.objects.get(period=OCT).status, Month.Status.OPEN)

    def test_left_block(self):
        run(sample_workbook())
        october = Month.objects.get(period=OCT)

        self.assertEqual(october.salary, D("50000.00"))
        self.assertEqual(october.salary_expression, "=1000+49000")
        self.assertEqual(october.total_expenses, D("72000.00"))
        self.assertEqual(october.notes, "-2000 лампа\n-500 книги")
        lines = {line.line_type.name: (line.amount, line.note) for line in october.lines.select_related("line_type")}
        self.assertEqual(
            lines,
            {
                "Rent": (D("15000.00"), ""),
                "Utility services": (D("1459.00"), "=1259+200"),
                "Additional incomes": (D("4600.00"), "3500 + 100€ подарунок; 1100 Петро + Іра"),
            },
        )

    def test_savings_and_crypto(self):
        run(sample_workbook())
        entries = [
            (e.amount, e.currency, e.destination.name if e.destination else None, e.raw_text, e.note)
            for e in SavingsEntry.objects.filter(month__period=OCT).select_related("destination").order_by("id")
        ]
        self.assertEqual(
            entries,
            [
                (D("100.00"), "USD", "моно", "$100 моно", ""),
                (D("6000.00"), "UAH", "облігації", "2000+4000 облігації", ""),
                (D("8000.00"), "UAH", None, "8000 + $200", ""),
                (D("200.00"), "USD", None, "8000 + $200", ""),
                (None, "UAH", None, "4000 (-1500)", ""),
                (D("2000.00"), "UAH", "Crypto", "2000 на USDT", "на USDT"),
            ],
        )
        old = SavingsEntry.objects.get(month__period=AUG)
        self.assertEqual((old.amount, old.raw_text), (D("2000.00"), "2000"))

    def test_allocations(self):
        run(sample_workbook())

        def allocations(period):
            month = Month.objects.get(period=period)
            return {
                a.category.name: (a.percent, a.limit_amount, a.actual, a.mark)
                for a in month.allocations.select_related("category")
            }

        self.assertEqual(
            allocations(OCT),
            {
                "Загальні Накопичення": (D("30.00"), D("15000.00"), D("12000.00"), "red"),
                "Їжа": (D("20.00"), D("10000.00"), None, ""),
                # 20% від «Як Є» Загальних Накопичень за 09.25.
                "Знято з накопичень": (D("20.00"), D("2000.00"), None, ""),
            },
        )
        # Частки 0.3 і 0.2 стали відсотками, «Як Є» Їжі має пріоритет над Food.
        self.assertEqual(
            allocations(SEP),
            {
                "Загальні Накопичення": (D("30.00"), D("12000.00"), D("10000.00"), "red"),
                "Їжа": (D("20.00"), D("8000.00"), D("8500.00"), "green"),
            },
        )
        # Старий місяць: Food стає «Як Є» Їжі.
        self.assertEqual(allocations(AUG), {"Їжа": (None, None, D("7000.00"), "")})

    def test_report_warnings(self):
        report = run(sample_workbook())
        messages = "\n".join(f"{label}: {message}" for label, message in report.warnings)
        self.assertIn("09.25: Їжа: «Як Є» 8 500.00 і Food 9 000.00 різні", messages)
        self.assertIn("10.25: Savings: «4000 (-1500)» складний", messages)
        self.assertIn("10.25: Additional incomes: 100 EUR", messages)
        self.assertNotIn("ліміт", messages)

    def test_limit_that_disagrees_with_rules_is_reported(self):
        book = sample_workbook()
        book.set("S6", 2500, "=N7*20%")
        book.set("O6", 14000)
        report = run(book)

        october = Month.objects.get(period=OCT)
        self.assertEqual(october.withdrawal_amount, D("2000.00"))
        self.assertIn(
            ("10.25", "Знято з накопичень: у таблиці ліміт 2 500.00, за правилами 2 000.00."),
            report.warnings,
        )
        # Число в «Ціль» без формули лишається ручним лімітом.
        self.assertEqual(october.savings_pool_allocation().limit_amount, D("14000.00"))

    def test_dry_run_saves_nothing(self):
        report = run(sample_workbook(), dry_run=True)
        self.assertEqual(len(report.created), 3)
        self.assertFalse(Month.objects.exists())
        self.assertIn("ПРОБНИЙ ЗАПУСК: зміни не збережено.", report.lines())

    def test_rerun_updates_open_month_and_skips_closed(self):
        run(sample_workbook())
        counts = (SavingsEntry.objects.count(), Month.objects.count())

        book = sample_workbook()
        book.set("C6", 80000)
        book.set("C7", 1)
        report = run(book)

        self.assertEqual(report.updated, ["10.25"])
        self.assertEqual(report.skipped_closed, ["08.25", "09.25"])
        self.assertEqual((SavingsEntry.objects.count(), Month.objects.count()), counts)
        self.assertEqual(Month.objects.get(period=OCT).total_expenses, D("80000.00"))
        self.assertIsNone(Month.objects.get(period=SEP).total_expenses)

    def test_overwrite_closed(self):
        run(sample_workbook())
        book = sample_workbook()
        book.set("C7", 1)
        report = run(book, overwrite_closed=True)

        september = Month.objects.get(period=SEP)
        self.assertEqual(september.total_expenses, D("1.00"))
        self.assertTrue(september.is_closed)
        self.assertEqual(report.updated, ["08.25", "09.25", "10.25"])


class ImportCommandTests(TestCase):
    def test_command_reads_file(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "budget.xlsx"
            sample_workbook().save_values(path)
            out = StringIO()
            call_command("import_excel", str(path), "--dry-run", stdout=out)
        self.assertIn("Створено місяців: 3", out.getvalue())
        self.assertFalse(Month.objects.exists())

    def test_unknown_category_is_an_error(self):
        book = sample_workbook()
        book.set("N3", "Невідома категорія")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "budget.xlsx"
            book.save_values(path)
            with self.assertRaisesMessage(CommandError, "Невідома категорія"):
                call_command("import_excel", str(path), stdout=StringIO())
